# -*- coding: utf-8 -*-
"""Weekly deck: one slide with the September 2026 rig schematic drawn for the report
(documentation/figures/report_rig_schematic.png). Appended at the end, and redo-safe: the
slide is located by its title ANYWHERE in the deck and rebuilt where it sits, so a rerun
cannot append a second copy if it has since been reordered; a duplicate left by an earlier
run is removed. Every other slide is proven identical and in the same order, all media
byte-identical, and the baseline is refreshed."""
import hashlib
import os
import shutil
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from wk_place_slides import ensure_page_number, prove_media_identical, prove_untouched

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
os.chdir(ROOT)
DECK = os.path.join("Documentation", "Decks", "Weekly progress updated.pptx")
BASE = os.path.join("Documentation", "Decks", ".local_baselines", "Weekly_progress.baseline.pptx")
FIG = os.path.join("documentation", "figures", "report_rig_schematic.png")
INK, GREY = RGBColor(0x25, 0x34, 0x39), RGBColor(0x5A, 0x63, 0x6B)
X0, X1 = 1.40, 12.07
TITLE = "RIG SCHEMATIC — THE PPD-UTM DIC RIG AS BUILT, SEPTEMBER 2026"


def is_schematic(slide):
    t = slide.shapes.title
    return t is not None and t.text.startswith("RIG SCHEMATIC")


def label(slide, x, y, w, h, text, size=11, colour=INK, bold=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0)
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = colour
    r.font.bold = bold
    r.font.name = "Calibri"
    return tb


def drop_slide(prs, idx):
    lst = prs.slides._sldIdLst
    sld = list(lst)[idx]
    rId = sld.rId
    lst.remove(sld)
    prs.part.drop_rel(rId)


def main():
    pre = os.path.join("Documentation", "Decks", ".local_baselines", "_pre_append.pptx")
    shutil.copyfile(DECK, pre)
    prs = Presentation(DECK)
    found = [i for i, sl in enumerate(prs.slides) if is_schematic(sl)]
    removed = found[1:]                    # keep the first; drop any duplicate an earlier run left
    for i in reversed(removed):
        drop_slide(prs, i)
    redo = bool(found)
    if redo:
        idx = found[0]
        s = prs.slides[idx]
        for sh in list(s.shapes):          # strip the generated content; placeholders stay
            if not sh.is_placeholder:
                sh._element.getparent().remove(sh._element)
    else:
        layout = next(l for l in prs.slide_layouts if l.name == "Two content light blue")
        s = prs.slides.add_slide(layout)
        idx = len(prs.slides) - 1
        for ph in list(s.placeholders):
            kind = str(ph.placeholder_format.type)
            if "TITLE" in kind:
                ph.text = TITLE
            else:
                ph._element.getparent().remove(ph._element)
    # python-pptx never clones a layout's slide-number placeholder, so it is copied in from a
    # neighbour on the same layout — a live slidenum field, not literal text.
    ensure_page_number(prs, s, idx + 1)
    from PIL import Image
    iw, ih = Image.open(FIG).size
    w = 7.2
    h = w * ih / iw
    s.shapes.add_picture(FIG, Inches(X0), Inches(2.1), Inches(w), Inches(h))
    x = X0 + w + 0.25
    label(s, x, 2.1, X1 - x, 0.3, "What it shows", size=11.5, bold=True)
    label(s, x, 2.45, X1 - x, 4.3,
          "Load path: two lead screws drive the crosshead; the S-beam load cell hangs from it, then "
          "the printed upper grip, the specimen with its two spray-paint markers, and the lower grip "
          "on the base holder.\n\n"
          "Strain channel: the Basler acA2440-35um + 25 mm lens on its printed mount post, ≈371 mm "
          "from the specimen; the LED strip above the field of view; the matte backdrop that keeps "
          "room light out.\n\n"
          "Data paths: force and position from the electronics box (ESP32, two TMC drivers, load-cell "
          "amplifier) over serial; frames over USB 3.0. Both E-stops cut motor power; the software "
          "stop and the 4.5 kN / 30 mm backstops sit on top.", size=9.5, colour=GREY)
    label(s, X0, 6.95, X1 - X0, 0.3,
          "Drawn for the project report (Chapter 3, rig schematic) by documentation/scripts/"
          "report_figs2.py · replaces the CAD-only schematic used earlier", size=9, colour=GREY)
    prs.save(DECK)

    # ---- proof: every other slide identical and in the same order, all media untouched.
    #      Position-based rather than part-name-based, because python-pptx renumbers the slide
    #      parts whenever the order changes (see wk_place_slides.py).
    kept = prove_untouched(pre, DECK, set(removed) | ({idx} if redo else set()), {idx})
    n_media = prove_media_identical(pre, DECK)
    prs2 = Presentation(DECK)
    mine = [i for i, sl in enumerate(prs2.slides) if is_schematic(sl)]
    assert mine == [idx], "expected exactly one schematic slide, at %d; found %s" % (idx, mine)
    print("schematic slide is %d of %d; removed duplicates: %s; %d other slides and %d media "
          "parts proven identical" % (idx + 1, len(prs2.slides), removed or "none", kept, n_media))
    for k in range(2):
        shutil.copyfile(DECK, BASE)
        assert hashlib.sha1(open(DECK, "rb").read()).hexdigest() == hashlib.sha1(open(BASE, "rb").read()).hexdigest()
    os.remove(pre)
    print("baseline refreshed:", hashlib.sha1(open(DECK, "rb").read()).hexdigest())


if __name__ == "__main__":
    main()
