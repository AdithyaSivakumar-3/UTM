# -*- coding: utf-8 -*-
"""Weekly deck: put the DIC good-practice slide where it belongs, and give both generated
slides the page number every other content slide has.

Two fixes, in this order, each proven and each idempotent:

  1. THE POSITION. The good-practice slide was appended at the end, but it belongs with the
     pre-study that produced it, so it is moved to sit immediately after "PRE-Study agenda".

  2. THE PAGE NUMBER. python-pptx's add_slide clones only the body placeholders of a layout —
     iter_cloneable_placeholders deliberately skips date, footer and slide-number — so the two
     generated slides never received the slide-number placeholder their layout offers, and
     showed no page number while 374 of their neighbours did. The placeholder is copied in
     whole from a donor slide on the SAME layout ("Two content light blue"), which keeps the
     position and the live <a:fld type="slidenum"> field; only the shape id and the field's
     cached text change. The number therefore follows the slide if it is ever moved again.
     The order of the two fixes matters: the cached text is written after the move, so it is
     already correct before PowerPoint recalculates the field.

A NOTE ON PROVING A MOVE. Part names cannot identify a slide across a reorder. python-pptx
renumbers the slide parts to match the slide order, and it does so IN MEMORY as soon as the
deck is opened -- so in a deck whose order no longer matches its stored names,
`slides[3].part.partname` says slide4.xml while the zip still holds the moved slide under
slide395.xml. Reading the archive by part name therefore reads the wrong slide. The proof
below is name-free: each slide is hashed from its own in-memory XML, canonicalised, and the
ORDERED list of hashes after the move must equal the list before it with exactly one element
moved. That catches any other reorder, insertion, deletion or edit to any slide.
"""
import copy
import hashlib
import io
import os
import shutil
import sys

from lxml import etree
from PIL import Image
from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
os.chdir(ROOT)
DECK = os.path.join("Documentation", "Decks", "Weekly progress updated.pptx")
BASE = os.path.join("Documentation", "Decks", ".local_baselines", "Weekly_progress.baseline.pptx")
TMP = os.path.join("Documentation", "Decks", ".local_baselines", "_pre_place.pptx")
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}

GENERATED = ("DIC GOOD-PRACTICE GUIDE", "RIG SCHEMATIC")
MOVE_TITLE = "DIC GOOD-PRACTICE GUIDE"
MOVE_AFTER = "PRE-Study agenda"


def title_of(slide):
    t = slide.shapes.title
    return (t.text or "").strip() if t is not None else ""


def number_ph(slide):
    for sh in slide.shapes:
        if sh.is_placeholder and "SLIDE_NUMBER" in str(sh.placeholder_format.type):
            return sh
    return None


def ensure_page_number(prs, slide, number):
    """Give `slide` the page number its layout offers, if it has none. Returns True if added.

    python-pptx never clones a layout's slide-number placeholder onto a new slide, so the
    placeholder is copied whole from a slide on the same layout that has one — position,
    formatting and the live <a:fld type="slidenum"> field included. Only the shape id and the
    field's cached text are changed, so the number still follows the slide when it is moved.
    Shared with wk_append_guide.py and wk_append_schematic.py so a redo cannot lose it again.
    """
    if number_ph(slide) is not None:
        return False
    donor = next(d for d in prs.slides
                 if d.slide_layout is slide.slide_layout and number_ph(d) is not None)
    sp = copy.deepcopy(number_ph(donor)._element)
    ext = sp.find("./p:nvSpPr/p:cNvPr/a:extLst", NS)           # drop the donor's creation id
    if ext is not None:
        ext.getparent().remove(ext)
    used = [int(e.get("id")) for e in slide._element.iter("{%s}cNvPr" % NS["p"])]
    sp.find("./p:nvSpPr/p:cNvPr", NS).set("id", str(max(used) + 1))
    fld = sp.find(".//a:fld", NS)
    assert fld is not None and fld.get("type") == "slidenum", "donor has no slidenum field"
    fld.find("./a:t", NS).text = str(number)       # cached text, until PowerPoint recalculates
    slide.shapes._spTree.append(sp)
    return True


def slide_hashes(path):
    """Canonical hash of each slide's XML, in DECK ORDER — never by part name (see above)."""
    prs = Presentation(path)
    out, titles = [], []
    for s in prs.slides:
        out.append(hashlib.sha1(etree.tostring(s._element, method="c14n")).hexdigest())
        titles.append(title_of(s))
    return out, titles


def prove_untouched(pre_path, post_path, pre_exclude, post_exclude):
    """Prove every slide but the named ones survived edit-for-edit, in the same order.

    Position-based, so it holds even when slides were dropped or moved and python-pptx has
    renumbered the parts. `pre_exclude` are indices in the deck before the edit (the slide
    being rewritten, plus any duplicate removed); `post_exclude` are indices after it.
    """
    h0, _ = slide_hashes(pre_path)
    h1, _ = slide_hashes(post_path)
    a = [h for i, h in enumerate(h0) if i not in set(pre_exclude)]
    b = [h for i, h in enumerate(h1) if i not in set(post_exclude)]
    assert a == b, ("slides other than the intended one differ: %d vs %d kept, first mismatch "
                    "at %s" % (len(a), len(b),
                               next((i for i in range(min(len(a), len(b))) if a[i] != b[i]), "—")))
    return len(a)


def prove_media_identical(pre_path, post_path):
    """Every media part the deck already had must still be byte-identical."""
    import zipfile
    zp, zd = zipfile.ZipFile(pre_path), zipfile.ZipFile(post_path)
    try:
        old = [n for n in zp.namelist() if n.startswith("ppt/media/")]
        missing = [n for n in old if n not in zd.namelist()]
        assert not missing, "media parts lost: %s" % missing[:5]
        bad = [n for n in old if zp.read(n) != zd.read(n)]
        assert not bad, "media parts rewritten: %s" % bad[:5]
        return len(old)
    finally:
        zp.close(); zd.close()


def check_pictures(path):
    prs = Presentation(path)
    bad = []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.shape_type == 13:
                try:
                    Image.open(io.BytesIO(sh.image.blob)).verify()
                except Exception:
                    bad.append(i)
    assert not bad, "broken pictures on slides %s" % sorted(set(bad))[:10]


def main():
    shutil.copyfile(DECK, TMP)
    h0, t0 = slide_hashes(DECK)

    # ------------------------------------------------------------------ 1. the position
    prs = Presentation(DECK)
    titles = [title_of(s) for s in prs.slides]
    src = next(i for i, t in enumerate(titles) if t.startswith(MOVE_TITLE))
    anchor = next(i for i, t in enumerate(titles) if t.startswith(MOVE_AFTER))
    dst = anchor + 1 if anchor < src else anchor
    if src != dst:
        ids = prs.slides._sldIdLst
        ids.insert(dst, list(ids)[src])       # lxml re-parents the element; no remove() needed
        prs.save(DECK)
        h1, t1 = slide_hashes(DECK)
        exp_h = [x for i, x in enumerate(h0) if i != src]
        exp_t = [x for i, x in enumerate(t0) if i != src]
        exp_h.insert(dst, h0[src])
        exp_t.insert(dst, t0[src])
        assert t1 == exp_t, "slide titles are not in the expected order"
        assert h1 == exp_h, "slide contents are not the old ones in the expected order"
        print("moved '%s' from slide %d to slide %d, after '%s'"
              % (MOVE_TITLE, src + 1, dst + 1, MOVE_AFTER))
        print("  proven: %d slides, every slide's content unchanged, order identical apart "
              "from that one move" % len(h1))
    else:
        h1, t1 = h0, t0
        print("'%s' already sits at slide %d, right after '%s'"
              % (MOVE_TITLE, src + 1, MOVE_AFTER))

    # ------------------------------------------------------------------ 2. the page number
    prs = Presentation(DECK)
    added = []
    for i, s in enumerate(prs.slides):
        if title_of(s).startswith(GENERATED) and ensure_page_number(prs, s, i + 1):
            added.append(i + 1)
    if added:
        prs.save(DECK)
    h2, t2 = slide_hashes(DECK)
    assert t2 == t1, "slide order changed while adding the page numbers"
    changed = [i + 1 for i in range(len(h2)) if h2[i] != h1[i]]
    assert changed == added, "changed slides %s, meant to change %s" % (changed, added)
    for i, s in enumerate(Presentation(DECK).slides):
        if title_of(s).startswith(GENERATED):
            ph = number_ph(s)
            assert ph is not None and ph._element.findall(".//a:fld", NS), \
                "slide %d still has no page-number field" % (i + 1)
            print("page number on slide %d ('%s…'): live slidenum field, cached %r"
                  % (i + 1, title_of(s)[:28], ph.text))
    print("page-number placeholder added to slide(s): %s" % (added or "none needed"))

    check_pictures(DECK)
    print("every picture in the deck still resolves")
    for _ in range(2):
        shutil.copyfile(DECK, BASE)
        assert hashlib.sha1(open(DECK, "rb").read()).hexdigest() == \
               hashlib.sha1(open(BASE, "rb").read()).hexdigest()
    os.remove(TMP)
    print("baseline refreshed:", hashlib.sha1(open(DECK, "rb").read()).hexdigest())


if __name__ == "__main__":
    main()
