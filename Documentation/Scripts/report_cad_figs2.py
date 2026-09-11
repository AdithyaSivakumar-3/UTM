# -*- coding: utf-8 -*-
"""Three more CAD figures for the report's design section, built from deck renders.

  cad_carriage_hybrid.png     the carriage with the chosen hybrid rib, arrowed and labelled
  cad_carriage_tslot.png      the carriage carrying the selected T-slot mushroom-head rail
  cad_housing_on_carriage.png the housing seated on the carriage, cropped out of the assembly

Sources are pictures already in the Weekly deck (slides 83, 90 and 96); this script reads the
BASELINE copy and writes only into documentation/figures/cad/.
"""
import io
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
os.chdir(ROOT)
DECK = os.path.join("Documentation", "Decks", ".local_baselines", "Weekly_progress.baseline.pptx")
OUT = os.path.join("documentation", "figures", "cad")
INK, ACC = "#253439", "#c0392b"
plt.rcParams["font.family"] = "Segoe UI"


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh


def pics(prs, n, min_in=1.5):
    out = []
    for sh in walk(prs.slides[n - 1].shapes):
        if sh.shape_type == 13:
            w, h = Emu(sh.width).inches, Emu(sh.height).inches
            if w >= min_in and h >= min_in:
                out.append(Image.open(io.BytesIO(sh.image.blob)).convert("RGB"))
    return out


def main():
    prs = Presentation(DECK)

    # ---- 1. the carriage with the hybrid rib, arrowed
    im = pics(prs, 83)[0]
    w, h = im.size
    fig, ax = plt.subplots(figsize=(w / 110, h / 110))
    ax.imshow(im)
    ax.set_xlim(0, w); ax.set_ylim(h, 0); ax.axis("off")
    # the rib is the diagonal web running from the back plate down-left to the platform
    ax.annotate("hybrid rib —\nstraight diagonal,\nfilleted at both ends",
                xy=(0.46 * w, 0.63 * h), xytext=(0.03 * w, 0.90 * h),
                fontsize=11.5, color=ACC, ha="left", va="center", fontweight="bold",
                arrowprops=dict(arrowstyle="-|>", color=ACC, lw=2.4, shrinkA=2, shrinkB=4))
    ax.annotate("camera platform", xy=(0.32 * w, 0.44 * h), xytext=(0.02 * w, 0.30 * h),
                fontsize=10.5, color=INK, ha="left", va="center",
                arrowprops=dict(arrowstyle="-", color="#8d949b", lw=1.1))
    ax.annotate("clamps the stand", xy=(0.70 * w, 0.72 * h), xytext=(0.72 * w, 0.95 * h),
                fontsize=10.5, color=INK, ha="center", va="center",
                arrowprops=dict(arrowstyle="-", color="#8d949b", lw=1.1))
    fig.subplots_adjust(0, 0, 1, 1)
    p = os.path.join(OUT, "cad_carriage_hybrid.png")
    fig.savefig(p, dpi=170, facecolor="white", bbox_inches="tight", pad_inches=0.02)
    plt.close(fig)
    print("  cad_carriage_hybrid.png")

    # ---- 2. the carriage carrying the selected T-slot mushroom-head rail
    im = pics(prs, 90)[1]
    im.save(os.path.join(OUT, "cad_carriage_tslot.png"))
    print("  cad_carriage_tslot.png  %dx%d" % im.size)

    # ---- 3. the housing seated on the carriage, cropped out of the side assembly
    im = pics(prs, 96)[1]
    w, h = im.size
    im.crop((int(0.24 * w), int(0.22 * h), int(0.78 * w), int(0.56 * h))) \
        .save(os.path.join(OUT, "cad_housing_on_carriage.png"))
    print("  cad_housing_on_carriage.png")


if __name__ == "__main__":
    main()
