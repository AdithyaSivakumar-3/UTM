"""Print the titles of named Weekly slides, using the TOPMOST wide box rule.

Never the longest text box: Weekly slides carry a body paragraph at top ~1.58 in and 10.6 in wide
that beats a 30-character title on any "longest" rule, which makes the slide unfindable.
"""
import sys

from pptx import Presentation
from pptx.util import Emu

DECK = sys.argv[1]
WANT = sys.argv[2:]


def title_of(s):
    best = None
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if not t or sh.width is None or sh.width < Emu(int(3.0 * 914400)):
            continue
        if sh.top is None:
            continue
        if best is None or sh.top < best[0]:
            best = (sh.top, t.split("\n")[0])
    return best[1] if best else "<no title>"


p = Presentation(DECK)
n = len(p.slides._sldIdLst)
print("%s\n  %d slides" % (DECK, n))
idxs = []
for w in WANT:
    if w == "end":
        idxs += list(range(n - 4, n + 1))
    elif ":" in w:
        a, b = w.split(":")
        idxs += list(range(int(a), int(b) + 1))
    else:
        idxs.append(int(w))
for i, s in enumerate(p.slides, 1):
    if i in idxs:
        print("  %4d  %s" % (i, title_of(s)[:88]))
