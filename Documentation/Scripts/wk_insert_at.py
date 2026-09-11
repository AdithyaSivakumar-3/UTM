"""Insert V6a slides into the Weekly deck at a named position. Nothing else is touched.

    python wk_insert_at.py <v6a.pptx> <weekly.pptx> 228:113 229:114

Each pair is  weeklyPositionAfterInsert : v6aSlideIndex , applied left to right.

THE GUARANTEE, checked before saving. Inserting shifts indices, so "every other slide is
byte-identical at the same index" is the wrong test. The right one: strip the fingerprints at the
inserted positions and what remains must equal the original list, in order, exactly. That proves
no existing slide was altered, dropped or reordered — only renumbered, which an insert must do.

Ordering is done through the sldId ELEMENTS, never index arithmetic.
"""
import copy
import hashlib
import io
import os
import re
import sys
import zipfile

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

SRC, DST = sys.argv[1], sys.argv[2]
PAIRS = [tuple(int(x) for x in a.split(":")) for a in sys.argv[3:]]
if not PAIRS:
    sys.exit("give at least one  weeklyPosition:v6aIndex  pair")

LAYOUT = "Two content light blue"
V6A_FIRST_PAGE = 141
S = 0.850
SRC_X0, SRC_Y0 = 0.40, 1.12
DST_X0, DST_Y0 = 1.40, 2.06
TITLE_SZ, MIN_PT = 2400, 6.0
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_EMBED, A_LINK = "{%s}embed" % R_NS, "{%s}link" % R_NS
A_ID = "{%s}id" % R_NS
PAGE_RE = re.compile(r"\bp(\d{3})\b")


def inch(v):
    return v / 914400.0


def fx(v):
    return Inches(DST_X0 + (inch(v) - SRC_X0) * S)


def fy(v):
    return Inches(DST_Y0 + (inch(v) - SRC_Y0) * S)


def flen(v):
    return Inches(inch(v) * S)


def fingerprint(slide):
    parts = []
    for sh in slide.shapes:
        parts.append("%s|%s|%s|%s|%s" % (sh.shape_type, sh.left, sh.top, sh.width, sh.height))
        if sh.has_text_frame:
            parts.append(sh.text_frame.text)
        if sh.has_table:
            parts += [c.text for r in sh.table.rows for c in r.cells]
    return hashlib.sha1("\x1f".join(parts).encode("utf-8", "replace")).hexdigest()


def is_chrome(sh):
    if sh.left is None or sh.top is None:
        return False
    l, t = inch(sh.left), inch(sh.top)
    w, h = inch(sh.width or 0), inch(sh.height or 0)
    txt = sh.text_frame.text.strip() if sh.has_text_frame else ""
    if "JÖNKÖPING" in txt.upper():
        return True
    if t < 0.60 and h <= 0.14 and w <= 0.70:
        return True
    if txt.isdigit() and l > 12.0 and t > 6.8:
        return True
    return False


def is_title(sh):
    return (sh.left is not None and sh.top is not None and sh.has_text_frame
            and inch(sh.top) < 0.75 and inch(sh.width or 0) > 11.0
            and sh.text_frame.text.strip())


def title_of(slide, house):
    best = None
    for sh in slide.shapes:
        if not sh.has_text_frame or sh.top is None:
            continue
        top, wid = inch(sh.top), inch(sh.width or 0)
        if top < (1.6 if house else 0.75) and wid > (9.0 if house else 11.0):
            if sh.text_frame.text.strip() and (best is None or top < best[0]):
                best = (top, sh.text_frame.text.strip().splitlines()[0])
    return best[1] if best else ""


def scale_runs(el, f):
    for rPr in el.iter():
        if not (rPr.tag.endswith("}rPr") or rPr.tag.endswith("}defRPr")
                or rPr.tag.endswith("}endParaRPr")):
            continue
        sz = rPr.get("sz")
        if sz:
            rPr.set("sz", str(max(int(MIN_PT * 100), int(round(int(sz) * f)))))


def furniture(slides):
    sldnum = house = None
    for s in slides:
        for sh in s.shapes:
            try:
                if sldnum is None and sh.is_placeholder and "SLIDE_NUMBER" in str(sh.placeholder_format.type):
                    sldnum = copy.deepcopy(sh._element)
            except Exception:
                pass
            if (house is None and sh.has_text_frame and sh.top is not None
                    and 1.0 < inch(sh.top) < 1.4 and inch(sh.width or 0) > 10
                    and 'cap="all"' in sh._element.xml):
                house = copy.deepcopy(sh._element)
        if sldnum is not None and house is not None:
            break
    return sldnum, house


def bring(dst, layout, srcs, sldnum, house):
    new = dst.slides.add_slide(layout)
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)
    if sldnum is not None:
        new.shapes._spTree.append(copy.deepcopy(sldnum))
    v6a_title = ""
    for sh in srcs.shapes:
        if is_title(sh):
            v6a_title = sh.text_frame.text.strip().splitlines()[0]
            continue
        if is_chrome(sh):
            continue
        el = copy.deepcopy(sh._element)
        new.shapes._spTree.append(el)
        # HYPERLINKS FIRST. a:hlinkClick carries an r:id into the SOURCE slide's relationship
        # part; copying the shape XML brings the id but not the relationship, so the new slide
        # points at an rId that does not exist. python-pptx does not validate that and will
        # happily save it — PowerPoint then refuses to open the file outright, with no clue
        # which slide is at fault. Found the hard way on the Otsu slide, the first one in this
        # deck to carry a link.
        for e in el.iter():
            if not (e.tag.endswith("}hlinkClick") or e.tag.endswith("}hlinkHover")):
                continue
            rid = e.get(A_ID)
            if not rid:
                continue
            try:
                url = srcs.part.rels[rid].target_ref
            except KeyError:
                del e.attrib[A_ID]              # a link with no target is dropped, not faked
                continue
            e.set(A_ID, new.part.relate_to(url, RT.HYPERLINK, is_external=True))
        for e in el.iter():
            for attr in (A_EMBED, A_LINK):
                rid = e.get(attr)
                if not rid:
                    continue
                try:
                    part = srcs.part.related_part(rid)
                except KeyError:
                    continue
                blob = getattr(part, "blob", None)
                if blob is None:
                    continue
                _, nrid = new.part.get_or_add_image_part(io.BytesIO(blob))
                e.set(attr, nrid)
    for sh in new.shapes:
        try:
            if sh.is_placeholder and "SLIDE_NUMBER" in str(sh.placeholder_format.type):
                continue
        except Exception:
            pass
        if sh.left is None or sh.top is None:
            continue
        sh.left, sh.top = fx(sh.left), fy(sh.top)
        if sh.width:
            sh.width = flen(sh.width)
        if sh.height:
            sh.height = flen(sh.height)
        scale_runs(sh._element, S)
        if sh.has_table:
            for c in sh.table.columns:
                c.width = flen(c.width)
            for r in sh.table.rows:
                r.height = flen(r.height)
    if v6a_title and house is not None:
        t_el = copy.deepcopy(house)
        new.shapes._spTree.append(t_el)
        new.shapes[-1].text_frame.text = v6a_title
        for dp in t_el.iter():
            if dp.tag.endswith("}defRPr") and dp.get("sz"):
                dp.set("sz", str(TITLE_SZ))
        for rp in t_el.iter():
            if rp.tag.endswith("}rPr") and rp.get("sz"):
                del rp.attrib["sz"]
    return v6a_title


def main():
    src, dst = Presentation(SRC), Presentation(DST)
    ss, ds = list(src.slides), list(dst.slides)
    before_n = len(ds)
    before_fp = [fingerprint(s) for s in ds]
    lst = dst.slides._sldIdLst
    layout = next(l for l in dst.slide_layouts if l.name == LAYOUT)
    sldnum, house = furniture(ds)

    for pos, v in PAIRS:
        if not (1 <= v <= len(ss)):
            sys.exit("V6a slide %d does not exist" % v)
        t = bring(dst, layout, ss[v - 1], sldnum, house)
        el = list(lst)[-1]
        lst.remove(el)
        lst.insert(pos - 1, el)
        print("inserted at Weekly %d   <- V6a %d   %r" % (pos, v, t[:62]))

    # ---- refs on the inserted slides, resolved against the deck as it now stands
    now = list(dst.slides)
    titles = {}
    for i, sl in enumerate(now):
        t = title_of(sl, True).strip().lower()
        if t:
            titles.setdefault(t, i + 1)

    def target(page):
        j = page - V6A_FIRST_PAGE
        if not (0 <= j < len(ss)):
            return None
        key = title_of(ss[j], False).strip().lower()
        if key in titles:
            return titles[key]
        hits = sorted({x for k, x in titles.items() if key and (key in k or k in key)})
        return hits[0] if len(hits) == 1 else None

    for pos, _v in PAIRS:
        for sh in now[pos - 1].shapes:
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if not PAGE_RE.search(run.text):
                        continue
                    out = run.text
                    for m in PAGE_RE.finditer(run.text):
                        pg = target(int(m.group(1)))
                        if pg:
                            out = out.replace(m.group(0), "p%d" % pg)
                            print("   ref %s -> p%d" % (m.group(0), pg))
                    run.text = out

    # ---- THE GUARANTEE: strip the inserted positions, the rest must match exactly and in order
    after_fp = [fingerprint(s) for s in list(dst.slides)]
    if len(after_fp) != before_n + len(PAIRS):
        sys.exit("ABORTED — expected %d slides, got %d" % (before_n + len(PAIRS), len(after_fp)))
    inserted = {p for p, _v in PAIRS}
    remains = [fp for i, fp in enumerate(after_fp, 1) if i not in inserted]
    if remains != before_fp:
        bad = [i + 1 for i, (a, b) in enumerate(zip(before_fp, remains)) if a != b]
        sys.exit("ABORTED — existing slides changed or reordered, first at original index %s"
                 % (bad[:5] or "?"))
    print("\nverified: %d inserted, all %d existing slides unchanged and in order"
          % (len(PAIRS), before_n))

    part = DST + ".part"
    dst.save(part)
    with zipfile.ZipFile(part) as z:
        if z.testzip() is not None:
            raise OSError("corrupt package")
        n = len([x for x in z.namelist() if x.startswith("ppt/slides/slide")])
    os.replace(part, DST)
    print("saved: %d slides, zip verified" % n)


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    main()
