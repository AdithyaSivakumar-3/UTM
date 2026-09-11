# -*- coding: utf-8 -*-
"""Weekly deck: one slide with the three links the project should carry.

  1. The initial reference this work started from — the department's own UTM repository.
  2. This project's own repository, the code and evidence behind every number in the deck.
  3. The VS Code introduction to code editing, the shortest way in for whoever picks it up.

The same three links go into the report (Section 1.1, Section 3.10 and Appendix E).

Redo-safe like the other wk_append_* scripts: the slide is located by its title anywhere in
the deck and rebuilt where it sits, a duplicate left by an earlier run is removed, every
other slide is proven identical and in the same order, all media byte-identical, and the
baseline is refreshed.
"""
import hashlib
import os
import shutil
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from wk_place_slides import ensure_page_number, prove_media_identical, prove_untouched

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
os.chdir(ROOT)
DECK = os.path.join("Documentation", "Decks", "Weekly progress updated.pptx")
BASE = os.path.join("Documentation", "Decks", ".local_baselines", "Weekly_progress.baseline.pptx")
INK, GREY = RGBColor(0x25, 0x34, 0x39), RGBColor(0x5A, 0x63, 0x6B)
LINK = RGBColor(0x1F, 0x6F, 0xB2)
X0, X1 = 1.40, 12.07
TITLE = "LINKS — WHERE THIS PROJECT CAME FROM, AND WHERE IT LIVES"

ROWS = [
    ("Where it came from",
     "https://github.com/cenmir/UTM",
     "The initial reference for this work was obtained from this repository: the department's "
     "own UTM — the frame, the electronics and the first control software. Everything in this "
     "deck was built on top of that starting point."),
    ("Where it lives now",
     "https://github.com/AdithyaSivakumar-3/UTM",
     "This project's own repository: the PyQt6 application and its twenty smart features, the "
     "98 tests, the evidence scripts that compute every figure in these slides, and the raw "
     "run data they read."),
    ("How to pick it up",
     "https://code.visualstudio.com/docs/introvideos/codeediting",
     "The code was written in Visual Studio Code. The editor's own introduction to code "
     "editing is the shortest way in for whoever continues the work."),
]


def is_links(slide):
    t = slide.shapes.title
    return t is not None and t.text.startswith("LINKS — WHERE THIS PROJECT")


def label(slide, x, y, w, h, text, size=11, colour=INK, bold=False, url=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0)
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = colour
    r.font.bold = bold
    r.font.name = "Calibri"
    if url:
        r.hyperlink.address = url          # clickable in the deck, like the other link slides
        r.font.underline = True
    return tb


def drop_slide(prs, idx):
    lst = prs.slides._sldIdLst
    sld = list(lst)[idx]
    rId = sld.rId
    lst.remove(sld)
    prs.part.drop_rel(rId)


def main():
    pre = os.path.join("Documentation", "Decks", ".local_baselines", "_pre_links.pptx")
    shutil.copyfile(DECK, pre)
    prs = Presentation(DECK)
    found = [i for i, s in enumerate(prs.slides) if is_links(s)]
    removed = found[1:]
    for i in reversed(removed):
        drop_slide(prs, i)
    redo = bool(found)
    if redo:
        idx = found[0]
        s = prs.slides[idx]
        for sh in list(s.shapes):
            if not sh.is_placeholder:
                sh._element.getparent().remove(sh._element)
    else:
        layout = next(l for l in prs.slide_layouts if l.name == "Two content light blue")
        s = prs.slides.add_slide(layout)
        idx = len(prs.slides) - 1
        for ph in list(s.placeholders):
            if "TITLE" in str(ph.placeholder_format.type):
                ph.text = TITLE
            else:
                ph._element.getparent().remove(ph._element)
    ensure_page_number(prs, s, idx + 1)

    label(s, X0, 2.34, X1 - X0, 0.3,
          "Three links this project should carry — the same three are in the report "
          "(Section 1.1, Section 3.10 and Appendix E).", size=11, colour=GREY)

    y = 2.90
    for head, url, text in ROWS:
        label(s, X0, y, 3.05, 0.3, head, size=13, colour=INK, bold=True)
        label(s, X0 + 3.20, y, X1 - X0 - 3.20, 0.3, url, size=12, colour=LINK, url=url)
        label(s, X0 + 3.20, y + 0.36, X1 - X0 - 3.20, 0.9, text, size=10.5, colour=GREY)
        y += 1.32

    label(s, X0, 6.90, X1 - X0, 0.3,
          "The upstream repository is credited as the starting point; this deck and the "
          "report cite both.", size=9, colour=GREY)
    prs.save(DECK)

    # ---- proof
    kept = prove_untouched(pre, DECK, set(removed) | ({idx} if redo else set()), {idx})
    n_media = prove_media_identical(pre, DECK)
    prs2 = Presentation(DECK)
    mine = [i for i, sl in enumerate(prs2.slides) if is_links(sl)]
    assert mine == [idx], "expected exactly one links slide, at %d; found %s" % (idx, mine)
    links = [r.hyperlink.address for sh in prs2.slides[idx].shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink.address]
    assert sorted(links) == sorted(u for _, u, _ in ROWS), links
    print("links slide is %d of %d; removed duplicates: %s; %d other slides and %d media parts "
          "proven identical" % (idx + 1, len(prs2.slides), removed or "none", kept, n_media))
    for u in links:
        print("   clickable: %s" % u)
    for _ in range(2):
        shutil.copyfile(DECK, BASE)
        assert hashlib.sha1(open(DECK, "rb").read()).hexdigest() == \
               hashlib.sha1(open(BASE, "rb").read()).hexdigest()
    os.remove(pre)
    print("baseline refreshed:", hashlib.sha1(open(DECK, "rb").read()).hexdigest())


if __name__ == "__main__":
    main()
