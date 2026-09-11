# -*- coding: utf-8 -*-
"""Weekly deck: append one slide mapping the iDICs good-practices guide onto the rig.

The pre-study deck (documentation/decks/Learnings.pptx, May 2026) digested the 166-page
guide into ten sections. This slide closes that loop at the end of the project: every rig
requirement the guide states, what the rig actually does about it, and the pattern rules the
two-marker reduction removes — which are exactly the rules the 2D-DIC future work inherits.

Redo-safe: the slide is located by its title anywhere in the deck (it now sits beside the
pre-study, not at the end — see wk_place_slides.py), a duplicate left by an earlier run is
removed, every other slide is proven identical and in the same order, and the baseline is
refreshed afterwards.
"""
import hashlib
import os
import shutil
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from wk_place_slides import ensure_page_number, prove_media_identical, prove_untouched

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
os.chdir(ROOT)
DECK = os.path.join("Documentation", "Decks", "Weekly progress updated.pptx")
BASE = os.path.join("Documentation", "Decks", ".local_baselines", "Weekly_progress.baseline.pptx")
INK, GREY = RGBColor(0x25, 0x34, 0x39), RGBColor(0x5A, 0x63, 0x6B)
GREEN, BLUE = RGBColor(0x1E, 0x84, 0x49), RGBColor(0x1F, 0x6F, 0xB2)
X0, X1 = 1.40, 12.07
TITLE = "DIC GOOD-PRACTICE GUIDE — WHAT IT ASKS FOR, AND WHAT THE RIG DOES"

ROWS = [
    ("", "What the iDICs guide asks", "On the PPD-UTM DIC rig"),
    ("RIG", "Global shutter, monochrome sensor",
     "Basler acA2440-35um — Sony IMX264 global shutter, mono"),
    ("RIG", "Every automatic function disabled",
     "Exposure, gain and gamma fixed in the material recipe; nothing auto"),
    ("RIG", "Fixed focal length, lockable focus and aperture rings",
     "Azure-2514MML 25 mm f/1.4; both rings set once at ≈371 mm and locked"),
    ("RIG", "Mount rigid AND decoupled from the load frame",
     "Printed carriage on an Al stand, on an 8 mm MDF plate, on a 2–3 mm TPU pad"),
    ("RIG", "Diffuse, symmetric lighting, locked before the test",
     "LED strip on printed uprights inside a matte enclosure; room light shut out"),
    ("RIG", "Motion blur under 0.5 px",
     "50 000 µs exposure at 1–2 mm/min — the specimen moves µm per frame"),
    ("RIG", "Image in the 50–80 % grey range",
     "Gamma 0.5 lifts the dark tones; the preset threshold sits in the histogram valley"),
    ("RIG", "Frame rate from the test speed (sampling theorem)",
     "≈20 fps delivered — oversamples these pull rates by a wide margin"),
    ("RIG", "Report every acquisition and analysis parameter",
     "Full setup written into every run's CSV header, plus the test registry (SF14)"),
    ("PATTERN", "Speckle 3–5 px per feature, random, matte, ≈50 % coverage",
     "No speckle — two sprayed markers of ≈10 000 px each"),
    ("PATTERN", "Subset 31–41 px, containing 3–5 speckles",
     "Not applicable — the marker is segmented and its centroid taken, not correlated"),
    ("PATTERN", "Step size 1/3 to 1/5 of the subset", "Not applicable"),
    ("PATTERN", "Virtual strain gauge = the length strain is reported over",
     "The marker separation itself: 80 mm or 45 mm, fixed when the dots are painted"),
]


def is_guide(slide):
    t = slide.shapes.title
    return t is not None and t.text.startswith("DIC GOOD-PRACTICE GUIDE")


def label(slide, x, y, w, h, text, size=11, colour=INK, bold=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0)
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = colour
    r.font.bold = bold
    r.font.name = "Calibri"
    return tb


def drop_slide(prs, idx):
    lst = prs.slides._sldIdLst
    sld = list(lst)[idx]
    rId = sld.rId
    lst.remove(sld)
    prs.part.drop_rel(rId)


def main():
    pre = os.path.join("Documentation", "Decks", ".local_baselines", "_pre_guide.pptx")
    shutil.copyfile(DECK, pre)
    prs = Presentation(DECK)
    # The slide is found by its title ANYWHERE in the deck, not only at the end: it has since
    # been moved next to the pre-study (wk_place_slides.py), and a rerun must rebuild it where
    # it now sits rather than append a second copy.
    found = [i for i, s in enumerate(prs.slides) if is_guide(s)]
    removed = found[1:]                     # keep the first; drop any duplicate an earlier run left
    for i in reversed(removed):
        drop_slide(prs, i)
    redo = bool(found)
    if redo:
        idx = found[0]
        s = prs.slides[idx]
        for sh in list(s.shapes):           # strip the generated content; placeholders stay
            if not sh.is_placeholder:
                sh._element.getparent().remove(sh._element)
    else:
        layout = next(l for l in prs.slide_layouts if l.name == "Two content light blue")
        s = prs.slides.add_slide(layout)
        idx = len(prs.slides) - 1
        for ph in list(s.placeholders):
            kind = str(ph.placeholder_format.type)
            if "TITLE" in kind:
                ph.text = TITLE
            else:
                ph._element.getparent().remove(ph._element)
    # python-pptx never clones a layout's slide-number placeholder, so it is copied in from a
    # neighbour on the same layout — a live slidenum field, not literal text.
    ensure_page_number(prs, s, idx + 1)

    label(s, X0, 2.34, X1 - X0, 0.3,
          "Read at the start of the project as a specification for the rig, not as background. "
          "The RIG rows are requirements this channel had to meet; the PATTERN rows are the ones "
          "the two-marker reduction removes — and the ones a 2D-DIC channel inherits.",
          size=10.5, colour=GREY)

    tb = s.shapes.add_table(len(ROWS), 3, Inches(X0), Inches(2.72),
                            Inches(X1 - X0), Inches(3.7)).table
    tb.columns[0].width = Inches(1.05)
    tb.columns[1].width = Inches(4.25)
    tb.columns[2].width = Inches(X1 - X0 - 5.30)
    for i, row in enumerate(ROWS):
        tb.rows[i].height = Inches(0.255)
        for j, v in enumerate(row):
            cell = tb.cell(i, j)
            cell.text = ""
            par = cell.text_frame.paragraphs[0]
            par.alignment = PP_ALIGN.CENTER if j == 0 else PP_ALIGN.LEFT
            run = par.add_run()
            run.text = v
            run.font.size = Pt(9.5)
            run.font.name = "Calibri"
            run.font.bold = (i == 0) or (j == 0)
            run.font.color.rgb = (INK if i == 0 else
                                  (GREEN if v == "RIG" else BLUE if v == "PATTERN" else INK))

    label(s, X0, 6.56, X1 - X0, 0.3,
          "The guide also names the path this project took: its recommended entry point is a flat "
          "tensile specimen, a single camera and a virtual extensometer — then full 2D-DIC. That is "
          "exactly this channel, and exactly the order of the future work.", size=10.5, colour=INK)
    label(s, X0, 6.98, X1 - X0, 0.3,
          "Source: iDICs, A Good Practices Guide for Digital Image Correlation, Edition 2 (October "
          "2025), digested in "
          "the pre-study deck documentation/decks/Learnings.pptx (May 2026), slides 7–8",
          size=9, colour=GREY)
    prs.save(DECK)

    # ---- proof: every other slide identical and in the same order, all media untouched.
    #      Position-based rather than part-name-based, because python-pptx renumbers the slide
    #      parts whenever the order changes (see wk_place_slides.py).
    kept = prove_untouched(pre, DECK, set(removed) | ({idx} if redo else set()), {idx})
    n_media = prove_media_identical(pre, DECK)
    prs2 = Presentation(DECK)
    mine = [i for i, sl in enumerate(prs2.slides) if is_guide(sl)]
    assert mine == [idx], "expected exactly one guide slide, at %d; found %s" % (idx, mine)
    print("guide slide is %d of %d; removed duplicates: %s; %d other slides and %d media parts "
          "proven identical" % (idx + 1, len(prs2.slides), removed or "none", kept, n_media))
    for _ in range(2):
        shutil.copyfile(DECK, BASE)
        assert hashlib.sha1(open(DECK, "rb").read()).hexdigest() == \
               hashlib.sha1(open(BASE, "rb").read()).hexdigest()
    os.remove(pre)
    print("baseline refreshed:", hashlib.sha1(open(DECK, "rb").read()).hexdigest())


if __name__ == "__main__":
    main()
